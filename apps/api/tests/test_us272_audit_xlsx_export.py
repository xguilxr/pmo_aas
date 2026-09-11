"""US-272 (FASE-9, revamp v2, pasada 1) — el export de auditoría pasa de
CSV a XLSX, mismo formato que el resto de las descargas del producto.
BUG-080 (columna `details`, el contexto del evento: job_id, errores,
conteos) sigue cubierto: la columna se conserva en la hoja.
"""
from __future__ import annotations

from datetime import UTC, datetime
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.models.audit import AuditLog
from tests.factories import create_admin_role, create_tenant, create_user, login


@pytest.mark.asyncio
async def test_export_xlsx_incluye_details(client, db_session):
    t = await create_tenant(db_session, slug="us272", name="US272")
    admin_role = await create_admin_role(db_session, t)
    await create_user(
        db_session, tenant=t, username="admin272",
        email="admin272@acme.example.com", password="Str0ng-Admin-1!",
        roles=[admin_role],
    )
    auth = await login(client, "admin272", "Str0ng-Admin-1!")

    db_session.add(AuditLog(
        tenant_id=str(t.id), action="ai.minute.generate", module="minutes",
        entity_type="ai_job", entity_id="job-123",
        details={"job_id": "job-123", "error": "boom", "count": 3},
        occurred_at=datetime.now(UTC),
    ))
    await db_session.commit()

    r = await client.get(
        "/api/v1/admin/audit-logs/export.xlsx", headers=auth["_authz"]
    )
    assert r.status_code == 200, r.text
    assert r.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "auditoria-us272-" in r.headers["content-disposition"]

    wb = load_workbook(BytesIO(r.content))
    ws = wb["Auditoría"]
    header = [c.value for c in ws[1]]
    assert "details" in header
    fila = [c.value for c in ws[2]]
    fila_str = " ".join(str(v) for v in fila)
    assert "job-123" in fila_str
    assert "boom" in fila_str
