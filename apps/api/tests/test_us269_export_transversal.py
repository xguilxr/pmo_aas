"""US-269 (FASE-8, revamp v2) — export XLSX de RAID y Cambios de todos los
proyectos filtrados, para el botón "Descargar Excel" de las pestañas
RAID y Cambios de `/pmo/reports`.

Dos proyectos en organizaciones distintas: el export filtrado por una
organización trae solo las filas de sus proyectos. El archivo abre con
openpyxl y la primera hoja tiene la columna "Proyecto".
"""
from __future__ import annotations

from io import BytesIO

import pytest
from openpyxl import load_workbook

from tests.factories import create_admin_role, create_tenant, create_user, login


async def _proyecto(client, auth, org_name: str, me_id: str) -> tuple[str, str]:
    org = (
        await client.post("/api/v1/organizations", json={"name": org_name}, headers=auth["_authz"])
    ).json()
    proj = (
        await client.post(
            "/api/v1/projects",
            json={
                "name": f"Proyecto {org_name}", "description": "d", "type": "innovacion",
                "priority": 3, "organization_id": org["id"], "pm_id": me_id,
            },
            headers=auth["_authz"],
        )
    ).json()
    return org["id"], proj["id"]


async def _area(client, auth, proj_id: str, org_id: str) -> str:
    area = (
        await client.post(
            "/api/v1/areas",
            json={"name": "Área", "organization_id": org_id},
            headers=auth["_authz"],
        )
    ).json()
    await client.put(
        f"/api/v1/admin/areas/{area['id']}/assignments",
        json={"scopes": [{"project_id": proj_id}]},
        headers=auth["_authz"],
    )
    return area["id"]


async def _escenario(client, db_session):
    t = await create_tenant(db_session, slug="us269", name="US269")
    admin_role = await create_admin_role(db_session, t)
    await create_user(
        db_session, tenant=t, username="admin269", email="admin269@pmoaas.example.com",
        password="Str0ng-Admin-1!", roles=[admin_role],
    )
    auth = await login(client, "admin269", "Str0ng-Admin-1!")
    me = (await client.get("/api/v1/auth/me", headers=auth["_authz"])).json()

    org_a, proj_a = await _proyecto(client, auth, "Org A", me["id"])
    org_b, proj_b = await _proyecto(client, auth, "Org B", me["id"])
    area_a = await _area(client, auth, proj_a, org_a)
    area_b = await _area(client, auth, proj_b, org_b)

    await client.post(
        f"/api/v1/projects/{proj_a}/risks",
        json={"title": "Riesgo A", "probability": 4, "impact": 4, "area_id": area_a},
        headers=auth["_authz"],
    )
    await client.post(
        f"/api/v1/projects/{proj_b}/risks",
        json={"title": "Riesgo B", "probability": 2, "impact": 2, "area_id": area_b},
        headers=auth["_authz"],
    )
    await client.post(
        f"/api/v1/projects/{proj_a}/change-requests",
        json={"title": "Cambio A", "type": "scope"},
        headers=auth["_authz"],
    )
    await client.post(
        f"/api/v1/projects/{proj_b}/change-requests",
        json={"title": "Cambio B", "type": "cost"},
        headers=auth["_authz"],
    )
    return {"h": auth["_authz"], "org_a": org_a, "org_b": org_b}


@pytest.mark.asyncio
async def test_export_raid_filtra_por_organizacion(client, db_session):
    e = await _escenario(client, db_session)
    r = await client.get(
        f"/api/v1/tenant/raid/export?organization_id={e['org_a']}", headers=e["h"]
    )
    assert r.status_code == 200, r.text
    assert "raid-org-a" in r.headers["content-disposition"]
    wb = load_workbook(BytesIO(r.content))
    ws = wb["Riesgos"]
    headers = [c.value for c in ws[1]]
    assert headers[:2] == ["Proyecto (folio)", "Proyecto"]
    titulos = [row[3].value for row in ws.iter_rows(min_row=2)]
    assert "Riesgo A" in titulos
    assert "Riesgo B" not in titulos


@pytest.mark.asyncio
async def test_export_changes_filtra_por_organizacion(client, db_session):
    e = await _escenario(client, db_session)
    r = await client.get(
        f"/api/v1/tenant/change-requests/export?organization_id={e['org_b']}", headers=e["h"]
    )
    assert r.status_code == 200, r.text
    assert "cambios-org-b" in r.headers["content-disposition"]
    wb = load_workbook(BytesIO(r.content))
    ws = wb["Cambios"]
    headers = [c.value for c in ws[1]]
    assert headers[:2] == ["Proyecto (folio)", "Proyecto"]
    titulos = [row[3].value for row in ws.iter_rows(min_row=2)]
    assert "Cambio B" in titulos
    assert "Cambio A" not in titulos


@pytest.mark.asyncio
async def test_export_sin_filtro_trae_de_todas_las_organizaciones(client, db_session):
    e = await _escenario(client, db_session)
    r = await client.get("/api/v1/tenant/raid/export", headers=e["h"])
    assert r.status_code == 200, r.text
    assert "raid-todas" in r.headers["content-disposition"]
    wb = load_workbook(BytesIO(r.content))
    ws = wb["Riesgos"]
    titulos = {row[3].value for row in ws.iter_rows(min_row=2)}
    assert {"Riesgo A", "Riesgo B"} <= titulos
